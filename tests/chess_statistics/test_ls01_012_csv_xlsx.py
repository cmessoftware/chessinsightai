"""Tests for LS01-012 — CSV/XLSX sheet Jugar en Lichess."""

from __future__ import annotations

import csv
import json
import sys
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook

SRC = Path(__file__).resolve().parents[2] / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

from chess_statistics.db import StatisticsRepository, connect, init_schema  # noqa: E402
from chess_statistics.export import (  # noqa: E402
    EXPORT_COLUMNS,
    SHEET_NAME,
    ExcelStatisticsExporter,
)
from chess_statistics.import_games import GameImportService  # noqa: E402

FIXTURE_TWO = Path(__file__).resolve().parent / "fixtures" / "cmess4401_rapid_two_games.ndjson"
USER = "cmess4401"


def _load_line(path: Path, index: int = 0) -> dict:
    lines = [line for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]
    return json.loads(lines[index])


def _imported_repo(tmp_path: Path, both: bool = False) -> StatisticsRepository:
    conn = connect(tmp_path / "ls.sqlite")
    init_schema(conn)
    repo = StatisticsRepository(conn)
    service = GameImportService()
    service.import_game(repo, _load_line(FIXTURE_TWO, 0), USER, fallback_local=False)
    if both:
        service.import_game(repo, _load_line(FIXTURE_TWO, 1), USER, fallback_local=False)
    return repo


def test_csv_and_xlsx_headers_and_one_data_row(tmp_path: Path):
    repo = _imported_repo(tmp_path)
    csv_path = tmp_path / "out.csv"
    xlsx_path = tmp_path / "out.xlsx"
    ExcelStatisticsExporter().export(repo, csv_path=csv_path, xlsx_path=xlsx_path)

    with csv_path.open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.reader(handle)
        headers = next(reader)
        data = next(reader)
    assert tuple(headers) == EXPORT_COLUMNS
    assert data[0] == "2026-09-15"
    assert data[1] == "https://lichess.org/tOsxrK57"
    assert data[2] == "15+10"
    assert data[4] == "Robertqwe"
    assert data[5] == "G"
    precision = float(data[11])
    assert 0 <= precision <= 100
    assert data[16] in ("lichess", "stockfish_local", "")
    assert data[19] == "rapid"
    assert data[20] == ""

    workbook = load_workbook(xlsx_path)
    assert workbook.sheetnames == [SHEET_NAME]
    sheet = workbook[SHEET_NAME]
    assert [cell.value for cell in sheet[1]] == list(EXPORT_COLUMNS)
    assert sheet["B2"].value == "https://lichess.org/tOsxrK57"
    assert isinstance(sheet["L2"].value, (int, float))
    assert 0 <= float(sheet["L2"].value) <= 100
    assert sheet["T2"].value == "rapid"
    assert sheet["U2"].value in ("", None)
    names = ZipFile(xlsx_path).namelist()
    assert "xl/vbaProject.bin" not in names
    assert not any(name.startswith("xl/macrosheets") for name in names)


def test_export_last_n_keeps_most_recent_rows(tmp_path: Path):
    repo = _imported_repo(tmp_path, both=True)
    csv_path = tmp_path / "out.csv"
    xlsx_path = tmp_path / "out.xlsx"
    ExcelStatisticsExporter().export(
        repo, csv_path=csv_path, xlsx_path=xlsx_path, usuario=USER, last_n=1
    )
    with csv_path.open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.reader(handle)
        next(reader)
        rows = list(reader)
    assert len(rows) == 1
    assert rows[0][1] == "https://lichess.org/tOsxrK57"
