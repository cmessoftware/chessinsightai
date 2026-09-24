"""Tests for LS01-020 — layer C profile, Entrenamiento sheet, profile JSON."""

from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import load_workbook

SRC = Path(__file__).resolve().parents[2] / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

from chess_statistics.aggregates import AggregateQueryService  # noqa: E402
from chess_statistics.db import StatisticsRepository, connect, init_schema  # noqa: E402
from chess_statistics.export import (  # noqa: E402
    SHEET_NAME,
    TRAINING_SHEET_NAME,
    ExcelStatisticsExporter,
)
from chess_statistics.import_games import GameImportService  # noqa: E402
from chess_statistics.training_profile import (  # noqa: E402
    ALLOWED_USE_EXPLAIN,
    MAX_SESSION_CANDIDATES,
    MAX_WEAKNESS_FOCI,
    PROFILE_SCHEMA_VERSION,
    build_weaknesses,
)

FIXTURE_TWO = Path(__file__).resolve().parent / "fixtures" / "cmess4401_rapid_two_games.ndjson"
USER = "cmess4401"


def _load_line(path: Path, index: int = 0) -> dict:
    lines = [line for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]
    return json.loads(lines[index])


def _repo_both_games(tmp_path: Path) -> StatisticsRepository:
    conn = connect(tmp_path / "ls.sqlite")
    init_schema(conn)
    repo = StatisticsRepository(conn)
    service = GameImportService()
    service.import_game(repo, _load_line(FIXTURE_TWO, 0), USER, fallback_local=False)
    service.import_game(repo, _load_line(FIXTURE_TWO, 1), USER, fallback_local=False)
    return repo


def test_training_profile_has_three_foci_and_candidates(tmp_path: Path):
    repo = _repo_both_games(tmp_path)
    report = AggregateQueryService(repo).report(USER, training_only=True)
    profile = report["training_profile"]
    assert profile["layer"] == "C"
    assert profile["schema_version"] == PROFILE_SCHEMA_VERSION
    assert profile["provenance"]["excluded_speeds"] == ["blitz", "bullet"]
    assert 1 <= len(profile["weaknesses"]) <= MAX_WEAKNESS_FOCI
    assert profile["weaknesses"][0]["rank"] == 1
    candidates = profile["candidate_positions"]
    assert 1 <= len(candidates) <= MAX_SESSION_CANDIDATES
    for item in candidates:
        assert item["allowed_uses"] == [ALLOWED_USE_EXPLAIN]
        assert "puzzle" not in item["allowed_uses"]
        assert item["fen_before"]
        assert item["eval_loss_cp"] >= 150


def test_export_adds_entrenamiento_sheet_when_training(tmp_path: Path):
    repo = _repo_both_games(tmp_path)
    xlsx_path = tmp_path / "out.xlsx"
    ExcelStatisticsExporter().export(
        repo,
        csv_path=tmp_path / "out.csv",
        xlsx_path=xlsx_path,
        usuario=USER,
        training_only=True,
    )
    workbook = load_workbook(xlsx_path)
    assert workbook.sheetnames == [SHEET_NAME, TRAINING_SHEET_NAME]
    sheet = workbook[TRAINING_SHEET_NAME]
    assert sheet["A1"].value == "Focos de entrenamiento"
    assert sheet["A2"].value == "Prioridad"
    assert sheet["A4"].value in (1, 2, 3)


def test_weakness_builder_returns_up_to_three_phase_foci():
    events = [
        {"phase": "opening", "eval_loss_cp": 160, "fecha": "2026-09-18", "judgment": "Mistake"},
        {"phase": "middlegame", "eval_loss_cp": 200, "fecha": "2026-09-17", "judgment": "Mistake"},
        {"phase": "middlegame", "eval_loss_cp": 180, "fecha": "2026-09-16", "judgment": "Blunder"},
        {"phase": "endgame", "eval_loss_cp": 300, "fecha": "2026-09-15", "judgment": "Blunder"},
    ]
    foci = build_weaknesses(events, until="2026-09-18")
    assert len(foci) == MAX_WEAKNESS_FOCI
    assert {item["phase"] for item in foci} == {"opening", "middlegame", "endgame"}


def test_export_single_sheet_without_training(tmp_path: Path):
    repo = _repo_both_games(tmp_path)
    xlsx_path = tmp_path / "plain.xlsx"
    ExcelStatisticsExporter().export(
        repo,
        csv_path=tmp_path / "plain.csv",
        xlsx_path=xlsx_path,
        usuario=USER,
    )
    workbook = load_workbook(xlsx_path)
    assert workbook.sheetnames == [SHEET_NAME]
