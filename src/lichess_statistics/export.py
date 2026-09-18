"""LS01-012 — CSV + XLSX export (Excel / Google Sheets safe)."""

from __future__ import annotations

import csv
from pathlib import Path
from typing import Any, Iterable

from lichess_statistics.db import StatisticsRepository

SHEET_NAME = "Jugar en Lichess"
LICHESS_GAME_URL = "https://lichess.org/{lichess_id}"

EXPORT_COLUMNS: tuple[str, ...] = (
    "Fecha",
    "Partida",
    "Ritmo",
    "Duración",
    "Rival",
    "G/T/P",
    "Ranking inicial",
    "Imprecisiones",
    "Errores",
    "Errores graves",
    "Pérdida prom. cp",
    "Precisión",
    "Precisión apertura",
    "Precisión mediojuego",
    "Precisión final",
    "Ranking final",
    "Comentarios",
)

_PRECISION_HEADERS = {
    "Precisión",
    "Precisión apertura",
    "Precisión mediojuego",
    "Precisión final",
}


def partida_url(lichess_id: str | None) -> str | None:
    if not lichess_id:
        return None
    return LICHESS_GAME_URL.format(lichess_id=lichess_id)


def _as_number(value: Any) -> int | float | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return value
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    if number.is_integer():
        return int(number)
    return number


def export_row(record: dict[str, Any]) -> dict[str, Any]:
    """One spreadsheet row; precisions stay numeric 0–100 (or empty)."""
    return {
        "Fecha": record.get("fecha"),
        "Partida": partida_url(record.get("lichess_id")),
        "Ritmo": record.get("ritmo"),
        "Duración": _as_number(record.get("duracion_segundos")),
        "Rival": record.get("rival"),
        "G/T/P": record.get("resultado"),
        "Ranking inicial": _as_number(record.get("ranking_inicial")),
        "Imprecisiones": _as_number(record.get("imprecisiones")),
        "Errores": _as_number(record.get("errores")),
        "Errores graves": _as_number(record.get("errores_graves")),
        "Pérdida prom. cp": _as_number(record.get("perdida_promedio_cp")),
        "Precisión": _as_number(record.get("precision_general")),
        "Precisión apertura": _as_number(record.get("precision_apertura")),
        "Precisión mediojuego": _as_number(record.get("precision_medio_juego")),
        "Precisión final": _as_number(record.get("precision_final")),
        "Ranking final": _as_number(record.get("ranking_final")),
        "Comentarios": "",
    }


def rows_from_repository(
    repo: StatisticsRepository,
    *,
    usuario: str | None = None,
    since: str | None = None,
    until: str | None = None,
    ritmo: str | None = None,
    last_n: int | None = None,
) -> list[dict[str, Any]]:
    records = repo.list_games_with_stats(
        usuario=usuario,
        since=since,
        until=until,
        ritmo=ritmo,
    )
    if last_n is not None:
        if last_n < 1:
            raise ValueError("last_n must be >= 1")
        records = records[-last_n:]
    return [export_row(record) for record in records]


def _cell_values(row: dict[str, Any]) -> list[Any]:
    return [row.get(column) for column in EXPORT_COLUMNS]


def write_csv(path: str | Path, rows: Iterable[dict[str, Any]]) -> Path:
    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    with out.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(EXPORT_COLUMNS)
        for row in rows:
            writer.writerow(_cell_values(row))
    return out


def _set_cell(cell: Any, header: str, value: Any) -> None:
    from openpyxl.styles import numbers

    if value is None or value == "":
        cell.value = None if header != "Comentarios" else ""
        return
    if header in _PRECISION_HEADERS and isinstance(value, (int, float)):
        cell.value = float(value)
        cell.number_format = "0.00"
        return
    cell.value = value
    if isinstance(value, int):
        cell.number_format = numbers.FORMAT_NUMBER
    elif isinstance(value, float):
        cell.number_format = "0.00"


def write_xlsx(path: str | Path, rows: Iterable[dict[str, Any]]) -> Path:
    """Workbook with a single data sheet; no VBA / macros."""
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_NAME
    for index, header in enumerate(EXPORT_COLUMNS, start=1):
        sheet.cell(1, index, header)
    for row_index, row in enumerate(rows, start=2):
        for col_index, header in enumerate(EXPORT_COLUMNS, start=1):
            _set_cell(sheet.cell(row_index, col_index), header, row.get(header))
    for index in range(1, len(EXPORT_COLUMNS) + 1):
        sheet.column_dimensions[get_column_letter(index)].width = 18
    workbook.vba_archive = None
    workbook.save(out)
    return out


class ExcelStatisticsExporter:
    """Export independent of HTTP; SQL stays in the repository."""

    def export(
        self,
        repo: StatisticsRepository,
        *,
        csv_path: str | Path,
        xlsx_path: str | Path,
        usuario: str | None = None,
        since: str | None = None,
        until: str | None = None,
        ritmo: str | None = None,
        last_n: int | None = None,
    ) -> tuple[Path, Path]:
        rows = rows_from_repository(
            repo,
            usuario=usuario,
            since=since,
            until=until,
            ritmo=ritmo,
            last_n=last_n,
        )
        return write_csv(csv_path, rows), write_xlsx(xlsx_path, rows)
